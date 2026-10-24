#!/usr/bin/env python3
"""日本食品標準成分表(八訂)増補2023年のExcelを data/foods.json に変換する。

使い方(プロジェクト直下で実行):
    pip install openpyxl
    python tools/convert_foods.py data/成分表.xlsx
    python tools/convert_foods.py data/成分表.xlsx --extra calcium,iron,zinc
    python tools/convert_foods.py data/成分表.xlsx --list-tags   # 検出した成分識別子の確認のみ

Excelの「成分識別子」行(ENERC_KCAL など)から列を自動検出する。
数値は100gあたり(可食部)。欠損値などの扱いは README.md の「成分表の変換」を参照。
"""
import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl が必要です: pip install openpyxl")

# 出力する栄養素キー -> 成分識別子の候補(先頭から順に探す)
BASE_TAGS = {
    "kcal": ["ENERC_KCAL"],
    "protein": ["PROT-", "PROTCAA"],
    "fat": ["FAT-", "FATNLEA"],
    "carb": ["CHOCDF", "CHOAVLM"],
    "fiber": ["FIB-", "FIB-TG", "FIBTG"],
}
# --extra で追加できる栄養素(js/nutrients.js の key と同じ名前にする)
EXTRA_TAGS = {
    "calcium": ["CA"],
    "iron": ["FE"],
    "zinc": ["ZN"],
    "magnesium": ["MG"],
    "vitaminB6": ["VITB6A"],
    "vitaminD": ["VITD"],
}

stats = Counter()


def parse_value(raw):
    """成分表の表記を数値に変換する。

    空欄・'-'・'…'・'*'  -> 0 (missing)
    'Tr' (微量)          -> 0 (tr)
    '(0)'                -> 0 (estimated)
    '(1.2)' / '[1.2]'    -> 1.2 (estimated: 推定値・計算値の括弧を外す)
    数値                 -> そのまま
    解釈できない文字列   -> 0 (unparsable, 警告を出す)
    """
    if raw is None:
        stats["missing"] += 1
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    s = unicodedata.normalize("NFKC", str(raw)).strip()  # '…' は '...' になる
    if s == "" or s in {"-", "...", "*", "−", "ー"}:
        stats["missing"] += 1
        return 0.0
    bracketed = False
    m = re.fullmatch(r"[(\[]\s*(.+?)\s*[)\]]", s)
    if m:
        bracketed = True
        s = m.group(1)
    if s.lower() == "tr":  # 'Tr' も '(Tr)' も微量として0
        stats["tr"] += 1
        return 0.0
    if bracketed:
        stats["estimated"] += 1
    try:
        return float(s.replace(",", ""))
    except ValueError:
        stats["unparsable"] += 1
        print(f"  警告: 解釈できない値を0にしました: {raw!r}", file=sys.stderr)
        return 0.0


def clean(cell):
    return unicodedata.normalize("NFKC", str(cell)).strip() if cell is not None else ""


def find_tag_row(wb):
    """成分識別子の行を含むシートと行番号を探す。"""
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            if any(clean(c) == "ENERC_KCAL" for c in row):
                return ws, r
    return None, None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=Path("data/foods.json"))
    ap.add_argument("--extra", default="", help="追加する栄養素(カンマ区切り): " + ",".join(EXTRA_TAGS))
    ap.add_argument("--list-tags", action="store_true", help="検出した成分識別子を表示して終了")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    ws, tag_r = find_tag_row(wb)
    if ws is None:
        sys.exit("成分識別子の行(ENERC_KCAL)が見つかりません。八訂の成分表Excelか確認してください。")

    rows = list(ws.iter_rows(values_only=True))
    tag_row = [clean(c) for c in rows[tag_r - 1]]
    tag_cols = {t: i for i, t in enumerate(tag_row) if t}
    print(f"シート '{ws.title}' の {tag_r} 行目を成分識別子行として検出 ({len(tag_cols)}列)")
    if args.list_tags:
        print(", ".join(tag_cols))
        return

    extras = [e for e in args.extra.split(",") if e]
    unknown = [e for e in extras if e not in EXTRA_TAGS]
    if unknown:
        sys.exit(f"未対応の栄養素: {unknown}(対応: {list(EXTRA_TAGS)})")

    wanted = {**BASE_TAGS, **{e: EXTRA_TAGS[e] for e in extras}}
    cols = {}
    for key, cands in wanted.items():
        col = next((tag_cols[c] for c in cands if c in tag_cols), None)
        if col is None:
            sys.exit(f"'{key}' の成分識別子 {cands} が見つかりません。--list-tags で確認してください。")
        cols[key] = col
        print(f"  {key}: 列{col + 1} ({tag_row[col]})")

    # 食品番号・食品名の列(見出し行から探す。なければ B列・D列)
    id_col, name_col = 1, 3
    for row in rows[:tag_r]:
        for i, c in enumerate(row):
            if clean(c) == "食品番号":
                id_col = i
            elif clean(c) == "食品名":
                name_col = i
    print(f"  食品番号: 列{id_col + 1} / 食品名: 列{name_col + 1}")

    need = max(id_col, name_col, *cols.values()) + 1
    foods, seen = [], set()
    for row in rows[tag_r:]:
        row = tuple(row) + (None,) * (need - len(row))  # 末尾の空セルが省略された行を補う
        raw_id = row[id_col]
        if isinstance(raw_id, float) and raw_id.is_integer():
            raw_id = int(raw_id)  # 1001.0 -> 1001
        fid = clean(raw_id)
        if not re.fullmatch(r"\d{1,5}", fid):
            continue  # 見出し・注記などの行
        fid = fid.zfill(5)
        name = re.sub(r"\s+", " ", clean(row[name_col]))
        if not name or fid in seen:
            continue
        seen.add(fid)
        n = {k: round(parse_value(row[c]), 2) for k, c in cols.items()}
        foods.append({"id": fid, "name": name, "n": n})

    out = {
        "version": 1,
        "source": "日本食品標準成分表(八訂)増補2023年",
        "per": "100g",
        "nutrients": list(cols),
        "foods": foods,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    wb.close()
    print(f"\n{len(foods)} 食品を {args.output} に出力しました")
    print("値の内訳: " + ", ".join(f"{k}={v}" for k, v in sorted(stats.items())))

    # 列の取り違えなどに気づくための確認表示
    if len(foods) < 2000:
        print(f"警告: 食品数が少なすぎます({len(foods)}件。八訂は約2,500件)。シートや列の検出を確認してください。", file=sys.stderr)
    if tag_row[cols["carb"]] == "CHOAVLM":
        print("警告: 炭水化物に CHOCDF ではなく CHOAVLM(利用可能炭水化物)を使いました。", file=sys.stderr)
    print("\n確認用サンプル(Excelの値と見比べてください):")
    for f in foods:
        if "全卵" in f["name"] and "生" in f["name"]:
            print(f"  {f['name']}: {f['n']}")
            break
    for f in foods[:2]:
        print(f"  {f['name']}: {f['n']}")


if __name__ == "__main__":
    main()
